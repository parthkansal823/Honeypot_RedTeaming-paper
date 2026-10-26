"""
Builds the Word report (and the project documentation) from Markdown.

    python report/src/build_report.py            # report  -> report/Project_Report.docx
    python report/src/build_report.py --docs     # docs/   -> docs/Project_Documentation.docx

Then run report/src/finalize_word.ps1 on the .docx to turn the equation
paragraphs into native Word equations, refresh the table of contents, the
lists of figures/tables and every cross-reference, and export a PDF.

Markdown conventions (a small, documented superset):
  # 1 Introduction              chapter (new page, "CHAPTER 1" + title)
  # ABSTRACT                    unnumbered front/back-matter heading
  ## 1.1 Title / ### / ####     headings 2-4 (numbers are written in the text)
  ![Caption.](images/x.png){#fig:id landscape nonumber}
  Table: Caption. {#tab:id widths=3,5,4}   (line before a pipe table)
  ```algorithm {#alg:id} Title  ... ```    boxed pseudocode with a caption
  $$ UnicodeMath $$ {#eq:id}               numbered equation
  [@fig:id] [@tab:id] [@alg:id] [@eq:id]   cross-references (Word REF fields)
  > text                                    shaded call-out box
  \newpage                                  page break
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "report"
FONT = "Times New Roman"
MONO = "Consolas"

PORTRAIT_W, PORTRAIT_H = 15.5, 22.0      # usable figure box on A4 portrait (cm)
# Diagram labels are 11 px x 1.15 (render scale) = 9.5 pt at natural size; below this scale
# they would print smaller than about 6.5 pt, so such wide diagrams are turned sideways.
MIN_LEGIBLE_SCALE = 0.685

# ----------------------------------------------------------------- helpers


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_style_font(style, name=FONT, size=None, bold=None, italic=None, color="000000"):
    f = style.font
    f.name = name
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), name)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(k)) is not None:
            del rfonts.attrib[qn(k)]


def para_format(style, align=None, line=1.5, before=0, after=6, keep_next=False, indent_first=None):
    pf = style.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = keep_next
    if indent_first is not None:
        pf.first_line_indent = Cm(indent_first)


def add_field(paragraph, instr, result="", bold=None, size=None):
    """Insert a complex field {instr} showing `result` until Word updates it."""
    def run_with(el):
        r = paragraph.add_run()
        if bold is not None:
            r.bold = bold
        if size:
            r.font.size = Pt(size)
        r._r.append(el)
        return r
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); b.set(qn("w:dirty"), "true")
    run_with(b)
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = f" {instr} "
    run_with(t)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    run_with(s)
    r = paragraph.add_run(result)
    if bold is not None:
        r.bold = bold
    if size:
        r.font.size = Pt(size)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run_with(e)


_BM_ID = [100]


def bookmark_start(paragraph, name):
    _BM_ID[0] += 1
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(_BM_ID[0])); s.set(qn("w:name"), name)
    paragraph._p.append(s)
    return _BM_ID[0]


def bookmark_end(paragraph, bid):
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid))
    paragraph._p.append(e)


def add_hyperlink(paragraph, url, text, size=None):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                         is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    st = OxmlElement("w:rStyle"); st.set(qn("w:val"), "Hyperlink"); rpr.append(st)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text; r.append(t)
    h.append(r)
    paragraph._p.append(h)


def bm_name(ref_id):
    return "_Ref_" + re.sub(r"[^A-Za-z0-9]", "_", ref_id)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true"); trPr.append(el)


def set_cell_margins(table, top=50, bottom=50, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); mar.append(el)
    tblPr.append(mar)


def set_table_borders(table, color="808080", size=4, inside=True):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def paragraph_border(paragraph, color="4C78A8", left_only=True):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for edge in (["left"] if left_only else ["top", "left", "bottom", "right"]):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "18"); el.set(qn("w:space"), "8"); el.set(qn("w:color"), color)
        bdr.append(el)
    pPr.append(bdr)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EEF3F9")
    pPr.append(shd)


def tiny_mark(paragraph):
    """Shrink a paragraph mark to 1 pt so a field that fills its page exactly does not
    push an empty line (and then a blank page) onto the next page."""
    pPr = paragraph._p.get_or_add_pPr()
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "2")
    rPr.append(sz)
    pPr.append(rPr)
    paragraph.paragraph_format.space_after = Pt(0)


def page_number_format(section, fmt="decimal", start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    elif pg.get(qn("w:start")) is not None:
        del pg.attrib[qn("w:start")]


# ------------------------------------------------------------ the builder


class Builder:
    def __init__(self, meta, header_text, docs_mode=False):
        self.meta = meta
        self.docs_mode = docs_mode
        self.doc = Document()
        self.header_text = header_text
        self.chapter = ""        # "1", "2", "A1", ...
        self.in_main = False
        self.landscape_now = False
        self.refs = {}           # id -> "Figure 3.2" etc
        self.counters = {}
        self.num_restart = 0
        self.rotate_wide = getattr(Builder, "rotate_default", False)
        self._styles()
        self._page_setup(self.doc.sections[0])

    # ----------------------------------------------------------- styles
    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        set_style_font(normal, FONT, 12)
        para_format(normal, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 0, 6)
        for name, size, align, before, after in (("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 0, 18),
                                                  ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 12, 6),
                                                  ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT, 10, 4),
                                                  ("Heading 4", 12, WD_ALIGN_PARAGRAPH.LEFT, 8, 2)):
            s = st[name]
            set_style_font(s, FONT, size, bold=True, italic=(name == "Heading 4"))
            para_format(s, align, 1.5 if name != "Heading 1" else 1.15, before, after, keep_next=True)
            s.paragraph_format.page_break_before = name == "Heading 1"   # never leaves a blank page
        cap = st["Caption"]
        set_style_font(cap, FONT, 10, bold=True, italic=False)
        para_format(cap, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 4, 10, keep_next=False)

        def new(name, base="Normal", kind=WD_STYLE_TYPE.PARAGRAPH):
            s = st.add_style(name, kind)
            s.base_style = st[base]
            return s

        s = new("Chapter Number")
        set_style_font(s, FONT, 16, bold=True)
        para_format(s, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 0, 6, keep_next=True)
        s = new("Front Title")
        set_style_font(s, FONT, 16, bold=True)
        para_format(s, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 0, 18, keep_next=True)
        s.paragraph_format.page_break_before = True
        s = new("Figure Paragraph")
        para_format(s, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 6, 2, keep_next=True)
        s = new("Table Caption", "Caption")
        para_format(s, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 10, 4, keep_next=True)
        s = new("Table Text")
        set_style_font(s, FONT, 10)
        para_format(s, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0)
        s = new("Code Text")
        set_style_font(s, MONO, 8.5, color="1F1F1F")
        para_format(s, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0)
        s = new("Equation Source")
        set_style_font(s, "Cambria Math", 12)
        para_format(s, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 6, 6)
        s = new("Callout")
        set_style_font(s, FONT, 11.5)
        para_format(s, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.3, 6, 8)
        s.paragraph_format.left_indent = Cm(0.4)
        s.paragraph_format.right_indent = Cm(0.4)
        s = new("Reference Entry")
        set_style_font(s, FONT, 11)
        para_format(s, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 8)
        s.paragraph_format.left_indent = Cm(0.9)
        s.paragraph_format.first_line_indent = Cm(-0.9)
        s = new("Abstract Text")
        set_style_font(s, FONT, 14)
        para_format(s, WD_ALIGN_PARAGRAPH.JUSTIFY, 2.0, 0, 6)
        s = new("Certificate Text")
        set_style_font(s, FONT, 14)
        para_format(s, WD_ALIGN_PARAGRAPH.JUSTIFY, 2.0, 0, 6)
        s = new("Title Page")
        set_style_font(s, FONT, 14)
        para_format(s, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 0, 0)
        for lvl in (1, 2, 3):
            try:
                t = st[f"TOC {lvl}"]
            except KeyError:
                t = st.add_style(f"TOC {lvl}", WD_STYLE_TYPE.PARAGRAPH)
            set_style_font(t, FONT, 12, bold=(lvl == 1))
            para_format(t, WD_ALIGN_PARAGRAPH.LEFT, 1.5, 0, 0)
            t.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
        try:
            tof = st["Table of Figures"]
        except KeyError:
            tof = st.add_style("Table of Figures", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(tof, FONT, 12)
        para_format(tof, WD_ALIGN_PARAGRAPH.LEFT, 1.5, 0, 0)
        try:
            hl = st["Hyperlink"]
        except KeyError:
            hl = st.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        hl.font.color.rgb = RGBColor(0x1F, 0x4E, 0x99)
        hl.font.underline = True

    def _page_setup(self, sec, landscape=False):
        w, h = Cm(21.0), Cm(29.7)
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = h, w
            sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.5)
            sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.5)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = w, h
            sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
            sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
        sec.header_distance = Cm(1.25)
        sec.footer_distance = Cm(1.25)

    # --------------------------------------------------------- sections
    def footer_page_number(self, section):
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(p, "PAGE", "1", size=11)

    def header_title(self, section):
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(self.header_text)
        r.italic = True
        r.font.size = Pt(9)
        r.font.name = FONT
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bt = OxmlElement("w:bottom")
        bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "4"); bt.set(qn("w:space"), "4"); bt.set(qn("w:color"), "999999")
        bdr.append(bt)
        pPr.append(bdr)

    def new_section(self, landscape=False, start_main=False):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page_setup(sec, landscape)
        sec.different_first_page_header_footer = False
        if start_main:
            page_number_format(sec, "decimal", 1)
            self.header_title(sec)
            self.footer_page_number(sec)
        else:
            page_number_format(sec, "decimal" if self.in_main else "lowerRoman")
        self.landscape_now = landscape
        return sec

    # ------------------------------------------------------- inline text
    INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[@[a-z]+:[A-Za-z0-9_\-]+\]|\[[^\]]+\]\([^)]+\)|https?://[^\s)]+|\$[^$]+\$)")

    def add_inline(self, p, text, size=None, base_bold=False):
        text = text.replace("\\|", "|")
        pos = 0
        for m in self.INLINE.finditer(text):
            if m.start() > pos:
                self._cited(p, text[pos:m.start()], size=size, bold=base_bold)
            tok = m.group(0)
            if tok.startswith("**"):
                self._run(p, tok[2:-2], bold=True, size=size)
            elif tok.startswith("`"):
                self._run(p, tok[1:-1], mono=True, size=(size or 12) - 1.5)
            elif tok.startswith("[@"):
                rid = tok[2:-1]
                label = self.refs.get(rid, rid)
                if rid in self.refs and not rid.startswith("eq:"):
                    add_field(p, f"REF {bm_name(rid)} \\h", label, size=size, bold=base_bold or None)
                else:
                    self._run(p, label, size=size, bold=base_bold)
            elif tok.startswith("["):
                mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
                add_hyperlink(p, mm.group(2), mm.group(1), size=size)
            elif tok.startswith("http"):
                add_hyperlink(p, tok.rstrip(".,;"), tok.rstrip(".,;"), size=size)
                tail = tok[len(tok.rstrip(".,;")):]
                if tail:
                    self._run(p, tail, size=size)
            elif tok.startswith("$"):
                self._run(p, tok[1:-1], italic=True, size=size)
            else:  # *italic*
                self._run(p, tok[1:-1], italic=True, size=size, bold=base_bold)
            pos = m.end()
        if pos < len(text):
            self._cited(p, text[pos:], size=size, bold=base_bold)

    PLACEHOLDER = re.compile(r"(\[[^\]]*(?:UID|Name|MONTH|Department|Designation|Academic|submission)[^\]]*\])")

    def _run(self, p, text, bold=False, italic=False, mono=False, size=None):
        """Add a run; [placeholders] get their own yellow-highlighted run."""
        text = text.replace("\\*", "*")
        parts = self.PLACEHOLDER.split(text)
        r = None
        for k, part in enumerate(parts):
            if part or len(parts) == 1:
                r = self._run_one(p, part, bold, italic, mono, size)
                if k % 2 == 1:
                    r.font.highlight_color = 7  # yellow
        return r

    def _run_one(self, p, text, bold, italic, mono, size):
        r = p.add_run(text)
        if bold:
            r.bold = True
        if italic:
            r.italic = True
        if mono:
            r.font.name = MONO
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), MONO)
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
            r._r.get_or_add_rPr().append(shd)
        if size:
            r.font.size = Pt(size)
        return r

    # ------------------------------------------------------- numbering
    def next_num(self, kind):
        key = (kind, self.chapter)
        self.counters[key] = self.counters.get(key, 0) + 1
        return f"{self.chapter}.{self.counters[key]}" if self.chapter else str(self.counters[key])

    def caption(self, kind, rid, text, style="Caption", nonumber=False):
        label = {"fig": "Figure", "tab": "Table", "alg": "Algorithm"}[kind]
        p = self.doc.add_paragraph(style=style)
        if nonumber:
            self.add_inline(p, text, size=10, base_bold=True)
            return p
        num = self.refs[rid].split(" ", 1)[1]
        chap, n = num.rsplit(".", 1)
        bid = bookmark_start(p, bm_name(rid))
        r = p.add_run(f"{label} {chap}.")
        r.bold = True
        add_field(p, f"SEQ {label} \\r {n} \\* ARABIC", n, bold=True)
        bookmark_end(p, bid)
        r2 = p.add_run(": ")
        r2.bold = True
        self.add_inline(p, text, size=10, base_bold=True)
        return p

    # ---------------------------------------------------------- blocks
    def heading(self, level, text):
        if text.endswith("{landscape}"):  # legacy marker; every page is portrait now
            text = text[: -len("{landscape}")].strip()
        if level == 1:
            m = re.match(r"^(\d+)\s+(.*)$", text)
            if m:  # numbered chapter (Heading 1 starts a new page by itself)
                if not self.in_main:
                    self.in_main = True
                    self.new_section(start_main=True)
                self.chapter = m.group(1)
                p = self.doc.add_paragraph(style="Heading 1")
                if self.docs_mode:
                    p.add_run(f"{m.group(1)}.  {m.group(2)}")
                    return
                r = p.add_run(f"CHAPTER {m.group(1)}")
                r.add_break(WD_BREAK.LINE)
                p.add_run(m.group(2).upper())
                return
            if text.upper().startswith(("REFERENCES", "APPENDICES")):
                self.chapter = "A" if text.upper().startswith("APPENDICES") else ""
            self.doc.add_paragraph(text.upper(), style="Heading 1")
            return
        if level == 2:
            m = re.match(r"^Appendix\s+(\d+)", text)
            if m:
                self.chapter = f"A{m.group(1)}"
        p = self.doc.add_paragraph(style=f"Heading {level}")
        if level == 2 and text.startswith("Appendix") and not self._last_is_heading1():
            p.paragraph_format.page_break_before = True
        self.add_inline(p, text)

    def _last_is_heading1(self):
        return self.doc.paragraphs and self.doc.paragraphs[-1].style.name == "Heading 1"

    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def ensure_portrait(self):
        if self.landscape_now:
            self.new_section(False)

    def paragraph(self, text, style=None):
        self.ensure_portrait()
        p = self.doc.add_paragraph(style=style)
        self.add_inline(p, text)
        return p

    def callout(self, text):
        self.ensure_portrait()
        p = self.doc.add_paragraph(style="Callout")
        paragraph_border(p)
        self.add_inline(p, text, size=11.5)

    def bullet_list(self, items, ordered=False):
        self.ensure_portrait()
        style = "List Number" if ordered else "List Bullet"
        try:
            num_id = self._restart_numbering() if ordered else None
        except Exception:  # noqa: BLE001 - fall back to continuous numbering
            num_id = None
        for depth, text in items:
            p = self.doc.add_paragraph(style=style if depth == 0 else (style + " 2"))
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if num_id is not None and depth == 0:
                pPr = p._p.get_or_add_pPr()
                numPr = OxmlElement("w:numPr")
                il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0")
                ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num_id))
                numPr.append(il); numPr.append(ni)
                pPr.append(numPr)
            self.add_inline(p, text)

    def _restart_numbering(self):
        """New w:num that points at 'List Number''s abstract list and restarts at 1."""
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        style = self.doc.styles["List Number"]
        base_num = style.element.pPr.numPr.numId.val
        abstract = None
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == str(base_num):
                abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        new_id = max(ids) + 1
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), abstract); num.append(a)
        ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); ov.append(so); num.append(ov)
        numbering.append(num)
        return new_id

    def figure(self, path, caption, rid, attrs):
        img_path = (REPORT / path) if not self.docs_mode else (ROOT / "docs" / path)
        if not img_path.exists():
            img_path = REPORT / path
        nonumber = "nonumber" in attrs
        # Every page is A4 portrait. A very wide diagram whose labels would print
        # below ~6.5 pt at text width is turned 90 degrees on its portrait page
        # instead, so it can use the full page height.
        with Image.open(img_path) as im:
            px_w, px_h = im.size
            dpi = (im.info.get("dpi") or (220, 220))[0] or 220
            nat_w = px_w / dpi * 2.54
            nat_h = px_h / dpi * 2.54
            upright = min(PORTRAIT_W / nat_w, PORTRAIT_H / nat_h, 1.0)
            sideways = min(PORTRAIT_W / nat_h, (PORTRAIT_H - 1.5) / nat_w, 1.0)
            rotate = (self.rotate_wide and "landscape" in attrs and upright < MIN_LEGIBLE_SCALE
                      and sideways > upright * 1.25)
            if rotate:
                rot_dir = REPORT / "images" / "rotated"
                rot_dir.mkdir(exist_ok=True)
                img_path = rot_dir / img_path.name
                im.rotate(90, expand=True).save(img_path, dpi=(dpi, dpi))
                nat_w, nat_h = nat_h, nat_w
        m = re.search(r"width=([\d.]+)", attrs)
        target_w = float(m.group(1)) if m else nat_w
        max_h = PORTRAIT_H - (1.5 if rotate else 0)
        scale = min(target_w / nat_w, PORTRAIT_W / nat_w, max_h / nat_h, 1.0 if not m else 99)
        p = self.doc.add_paragraph(style="Figure Paragraph")
        run = p.add_run()
        shape = run.add_picture(str(img_path), width=Cm(nat_w * scale))
        docpr = shape._inline.docPr
        docpr.set("descr", re.sub(r"[*`\[\]@]", "", caption))
        docpr.set("title", Path(path).stem)
        self.caption("fig", rid, caption, nonumber=nonumber)

    def table(self, caption, rid, attrs, rows):
        if self.landscape_now:
            self.new_section(False)
        nonumber = "nonumber" in attrs
        if caption:
            self.caption("tab", rid, caption, style="Table Caption", nonumber=nonumber)
        header, body = rows[0], rows[1:]
        ncol = len(header)
        m = re.search(r"widths=([\d.,]+)", attrs)
        widths = [float(x) for x in m.group(1).split(",")] if m else [1.0] * ncol
        if len(widths) != ncol:
            widths = (widths + [1.0] * ncol)[:ncol]
        total = sum(widths)
        cm_w = [PORTRAIT_W * w / total for w in widths]
        t = self.doc.add_table(rows=1, cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_table_borders(t)
        set_cell_margins(t)
        for i, txt in enumerate(header):
            c = t.rows[0].cells[i]
            c.width = Cm(cm_w[i])
            set_cell_shading(c, "DCE6F1")
            cp = c.paragraphs[0]
            cp.style = self.doc.styles["Table Text"]
            self.add_inline(cp, txt.strip(), size=10, base_bold=True)
        set_repeat_header(t.rows[0])
        for k, row in enumerate(body):
            cells = t.add_row().cells
            set_cant_split(t.rows[-1])
            for i in range(ncol):
                txt = row[i].strip() if i < len(row) else ""
                c = cells[i]
                c.width = Cm(cm_w[i])
                if k % 2 == 1:
                    set_cell_shading(c, "F7F9FC")
                cp = c.paragraphs[0]
                cp.style = self.doc.styles["Table Text"]
                bold = txt.startswith("**") and txt.endswith("**")
                self.add_inline(cp, txt[2:-2] if bold else txt, size=10, base_bold=bold)
        if len(t.rows) <= 18:
            for row in t.rows[:-1]:
                for c in row.cells:
                    for cp in c.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 1.0

    def code_block(self, lines, rid=None, title=None, kind="code"):
        if self.landscape_now:
            self.new_section(False)
        if kind == "algorithm":
            self.caption("alg", rid, title, style="Table Caption")
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if kind == "algorithm":
            tblPr = t._tbl.tblPr
            b = OxmlElement("w:tblBorders")
            for edge, sz in (("top", 12), ("bottom", 12)):
                el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
                el.set(qn("w:space"), "0"); el.set(qn("w:color"), "333333"); b.append(el)
            tblPr.append(b)
        else:
            set_table_borders(t, "C8C8C8", 4, inside=False)
        set_cell_margins(t, 80, 80, 140, 140)
        c = t.rows[0].cells[0]
        c.width = Cm(PORTRAIT_W)
        set_cell_shading(c, "FAFAFA" if kind == "algorithm" else "F4F6F8")
        first = True
        for line in lines:
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            p.style = self.doc.styles["Code Text"]
            txt = line.replace("▷", "//") if kind == "algorithm" else line
            if kind == "algorithm" and re.match(r"^(Input|Output):", txt):
                k, v = txt.split(":", 1)
                self._run(p, k + ":", bold=True, mono=True, size=8.5)
                self._run(p, v, mono=True, size=8.5)
            else:
                r = p.add_run(txt if txt else " ")
                r.font.name = MONO
                r.font.size = Pt(8.5)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 1.0

    def equation(self, body, rid):
        if self.landscape_now:
            self.new_section(False)
        num = self.refs[rid].split("(")[1].rstrip(")")
        p = self.doc.add_paragraph(style="Equation Source")
        r = p.add_run(f"{body.strip()}#({num})")
        r.font.name = "Cambria Math"

    # --------------------------------------------------------- pre-pass
    def collect_refs(self, texts):
        chapter = ""
        counts = {}
        for text in texts:
            in_code = False
            for line in text.split("\n"):
                if line.strip().startswith("```"):
                    starts_alg = line.strip().startswith("```algorithm")
                    if in_code or not starts_alg:
                        in_code = not in_code
                        continue
                    in_code = True
                elif in_code:
                    continue
                h1 = re.match(r"^# (\d+)\s", line)
                if h1:
                    chapter = h1.group(1)
                elif line.startswith("# "):
                    chapter = "A" if "APPENDICES" in line.upper() else ""
                ap = re.match(r"^## Appendix\s+(\d+)", line)
                if ap:
                    chapter = f"A{ap.group(1)}"
                for kind, pat in (("fig", r"^!\[.*\]\(.*\)\{#(fig:[^\s}]+)([^}]*)\}"),
                                  ("tab", r"^Table:.*\{#(tab:[^\s}]+)([^}]*)\}"),
                                  ("alg", r"^```algorithm \{#(alg:[^\s}]+)\}"),
                                  ("eq", r"^\$\$.*\$\$\s*\{#(eq:[^\s}]+)\}")):
                    m = re.match(pat, line)
                    if m:
                        if kind in ("fig", "tab") and "nonumber" in (m.group(2) if m.lastindex >= 2 else ""):
                            continue
                        counts[(kind, chapter)] = counts.get((kind, chapter), 0) + 1
                        n = counts[(kind, chapter)]
                        num = f"{chapter}.{n}" if chapter else str(n)
                        self.refs[m.group(1)] = {"fig": f"Figure {num}", "tab": f"Table {num}",
                                                 "alg": f"Algorithm {num}", "eq": f"Equation ({num})"}[kind]

    # ------------------------------------------------------------ parse
    def render_markdown(self, text):
        lines = text.split("\n")
        i = 0
        para = []

        def flush():
            if para:
                self.paragraph(" ".join(s.strip() for s in para))
                para.clear()

        while i < len(lines):
            line = lines[i]
            s = line.strip()
            if s.startswith("<!--"):
                flush()
                while "-->" not in lines[i]:
                    i += 1
                i += 1
                continue
            if not s:
                flush(); i += 1; continue
            if s == "\\newpage":
                flush()
                if self.landscape_now:
                    self.new_section(False)
                else:
                    self.page_break()
                i += 1; continue
            m = re.match(r"^(#{1,4})\s+(.*)$", s)
            if m:
                flush(); self.heading(len(m.group(1)), m.group(2).strip()); i += 1; continue
            m = re.match(r"^!\[(.*)\]\((.*?)\)\{#(fig:[^\s}]+)([^}]*)\}", s)
            if m:
                flush(); self.figure(m.group(2), m.group(1), m.group(3), m.group(4)); i += 1; continue
            m = re.match(r"^Table:\s*(.*?)\s*\{#(tab:[^\s}]+)([^}]*)\}", s)
            if m:
                flush()
                cap, rid, attrs = m.group(1), m.group(2), m.group(3)
                i += 1
                while i < len(lines) and not lines[i].strip():
                    i += 1
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    raw = lines[i].strip().strip("|")
                    cells = re.split(r"(?<!\\)\|", raw)
                    if not all(re.fullmatch(r"\s*:?-{2,}:?\s*", c) for c in cells):
                        rows.append(cells)
                    i += 1
                self.table(cap, rid, attrs, rows)
                continue
            m = re.match(r"^```(\w+)?\s*(?:\{#(alg:[^\s}]+)\})?\s*(.*)$", s)
            if m:
                flush()
                kind = "algorithm" if m.group(1) == "algorithm" else "code"
                rid, title = m.group(2), m.group(3)
                i += 1
                block = []
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    block.append(lines[i].rstrip())
                    i += 1
                i += 1
                self.code_block(block, rid, title, kind)
                continue
            m = re.match(r"^\$\$(.*)\$\$\s*\{#(eq:[^\s}]+)\}", s)
            if m:
                flush(); self.equation(m.group(1), m.group(2)); i += 1; continue
            if s.startswith(">"):
                flush()
                block = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    block.append(lines[i].strip()[1:].strip())
                    i += 1
                self.callout(" ".join(block))
                continue
            m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
            if m and not s.startswith("**"):
                flush()
                ordered = m.group(2)[0].isdigit()
                items = []
                while i < len(lines):
                    mm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                    if not mm:
                        if lines[i].strip() and lines[i].startswith("   ") and items:
                            items[-1] = (items[-1][0], items[-1][1] + " " + lines[i].strip())
                            i += 1
                            continue
                        break
                    items.append((1 if len(mm.group(1)) >= 2 else 0, mm.group(3)))
                    i += 1
                if self._in_references:
                    for _, t in items:
                        self.reference(t)
                else:
                    self.bullet_list(items, ordered)
                continue
            para.append(line)
            i += 1
        flush()

    _in_references = False
    _ref_n = 0
    _cite_re = None

    def reference(self, text):
        self._ref_n += 1
        p = self.doc.add_paragraph(style="Reference Entry")
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
        bid = bookmark_start(p, f"_Cite_{self._ref_n}")
        self._run(p, f"{self._ref_n}.\t", size=11)
        self.add_inline(p, text, size=11)
        bookmark_end(p, bid)

    # ------------------------------------------------- in-text citations
    def load_citations(self, ref_md):
        """Index the reference list so in-text citations such as (Spitzner, 2003a) or
        Arp et al. (2022) can be turned into links to their entry."""
        self._cites = {}
        n = 0
        for line in ref_md.splitlines():
            m = re.match(r"^\d+\.\s+(.*?)\s*\((\d{4}[a-z]?|n\.d\.)\)", line)
            if not m:
                continue
            n += 1
            authors, year = m.group(1), m.group(2)
            first = authors.split(",")[0].strip()
            surnames = re.findall(r"([A-ZÀ-Ž][\w'\-À-ž]+(?: [A-ZÀ-Ž][\w'\-À-ž]+)?), [A-Z]\.", authors) or [first]
            initial = re.match(r"^[^,]+, ([A-Z]\.)", authors)
            many = len(surnames) + (5 if "et al" in authors else 0)
            self._cites.setdefault((first, year), []).append(
                {"n": n, "disc": surnames[1] if len(surnames) > 1 else None,
                 "initial": f"{first} {initial.group(1)}" if initial else None, "many": many})
        names = "|".join(re.escape(x) for x in sorted({k[0] for k in self._cites}, key=len, reverse=True))
        year = r"(?P<y>\d{4}[a-z]?|n\.d\.)"
        # one item inside a parenthetical citation: "Arp et al., 2022"
        self._cite_re = re.compile(r"(?P<s>" + names + r")\b(?P<mid>[^();\d]{0,60}?),?\s*" + year)
        # a narrative citation: "Arp et al. (2022)" -> "Arp et al. [3]"
        self._cite_narr = re.compile(r"\b(?P<s>" + names + r")\b(?P<mid>[^();\d]{0,60}?)\(" + year + r"\)")

    def _cite_target(self, m):
        cands = self._cites.get((m.group("s"), m.group("y")))
        if not cands:
            return None
        if len(cands) == 1:
            return cands[0]["n"]
        span = m.group(0)
        ranked = sorted(cands, key=lambda c: -c["many"])
        for c in ranked:  # an initial or a co-author named in the citation decides first
            if (c["initial"] and c["initial"] in span) or (c["disc"] and c["disc"] in span):
                return c["n"]
        for c in ranked:
            if "et al" in span and c["many"] >= 3:
                return c["n"]
        return min(cands, key=lambda c: c["many"])["n"]

    def _group_numbers(self, inner):
        """Reference numbers for the inside of "(Arp et al., 2022; Spitzner, 2003a)", or None
        if any item is not a citation (so ordinary brackets are left alone)."""
        nums, last = [], None
        for item in inner.split(";"):
            item = item.strip()
            m = self._cite_re.fullmatch(item)
            if m:
                n, last = self._cite_target(m), m.group("s")
            elif last and (last, item) in self._cites:   # "(Government of India, 2000; 2023)"
                n = self._cites[(last, item)][0]["n"]
            else:
                return None
            if n is None:
                return None
            nums.append(n)
        return nums

    def _cited(self, p, text, size=None, bold=False):
        """Plain text with numbered citations: "(Arp et al., 2022)" becomes "[3]" and
        "Arp et al. (2022)" becomes "Arp et al. [3]"; every number links to its reference."""
        if not self._cite_re or self._in_references:
            self._run(p, text, size=size, bold=bold)
            return
        spans = []  # (start, end, [numbers], text kept before the bracket)
        for m in re.finditer(r"\(([^()]*?(?:\d{4}[a-z]?|n\.d\.)[^()]*?)\)", text):
            nums = self._group_numbers(m.group(1))
            if nums:
                spans.append((m.start(), m.end(), nums))
        for m in self._cite_narr.finditer(text):
            year_start = m.start("y") - 1
            if any(a <= year_start < b for a, b, _ in spans):
                continue
            n = self._cite_target(m)
            if n is not None:
                spans.append((year_start, m.end(), [n]))
        pos = 0
        for a, b, nums in sorted(spans):
            if a < pos:
                continue
            self._run(p, text[pos:a], size=size, bold=bold)
            self._run(p, "[", size=size, bold=bold)
            for k, n in enumerate(sorted(set(nums))):
                if k:
                    self._run(p, ", ", size=size, bold=bold)
                r = self._run(p, str(n), size=size, bold=bold)
                h = OxmlElement("w:hyperlink")
                h.set(qn("w:anchor"), f"_Cite_{n}")
                h.set(qn("w:history"), "1")
                r._r.addprevious(h)
                h.append(r._r)
            self._run(p, "]", size=size, bold=bold)
            pos = b
        if pos < len(text):
            self._run(p, text[pos:], size=size, bold=bold)

    # -------------------------------------------------------- front pages
    def title_page(self):
        d, m = self.doc, self.meta
        sec = d.sections[0]
        sec.different_first_page_header_footer = True
        page_number_format(sec, "lowerRoman", 1)
        self.footer_page_number(sec)

        def line(text, size, bold=False, italic=False, space_after=0, line_sp=1.5):
            p = d.add_paragraph(style="Title Page")
            p.paragraph_format.space_after = Pt(space_after)
            p.paragraph_format.line_spacing = line_sp
            r = self._run(p, text, bold=bold, italic=italic, size=size)
            return p

        for _ in range(2):
            line("", 12)
        line(m["title"].upper(), 18, bold=True, space_after=30)
        line(m.get("report_type", "A PROJECT REPORT"), 14, bold=True, space_after=24)
        line("Submitted by", 14, bold=True, italic=True, space_after=12)
        for s in m["students"].split(";"):
            line(s.strip().upper(), 16, bold=True)
        line("", 12)
        line("in partial fulfillment for the award of the degree of", 14, bold=True, italic=True, space_after=12)
        line(m["degree"].upper(), 16, bold=True, space_after=10)
        line("IN", 14, space_after=10)
        line(m["branch"].upper(), 14, space_after=60)
        line(m["university"], 16, bold=True)
        line(m["month_year"].upper(), 14)
        # bonafide certificate (Front Title starts a new page)
        p = d.add_paragraph(style="Front Title")
        p.add_run("BONAFIDE CERTIFICATE")
        p = d.add_paragraph(style="Certificate Text")
        self._run(p, "Certified that this project report “", size=14)
        self._run(p, m["title"], bold=True, size=14)
        self._run(p, "” is the bonafide work of “", size=14)
        self._run(p, ", ".join(s.split("(")[0].strip() for s in m["students"].split(";")), bold=True, size=14)
        self._run(p, "” who carried out the project work under my/our supervision.", size=14)
        d.add_paragraph()
        t = d.add_table(rows=6, cols=2)
        t.autofit = False
        rows = [("SIGNATURE", "SIGNATURE"), (m["hod"], m["supervisor"]), ("HEAD OF THE DEPARTMENT", "SUPERVISOR"),
                (m["department"], m["supervisor_designation"]), ("", m["department"]), ("", "")]
        for r_i, (a, b) in enumerate(rows):
            for c_i, txt in enumerate((a, b)):
                c = t.rows[r_i].cells[c_i]
                c.width = Cm(7.75)
                cp = c.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                cp.paragraph_format.line_spacing = 1.3
                self._run(cp, txt, bold=(r_i in (0, 2)), size=12)
        for _ in range(2):
            d.add_paragraph()
        p = d.add_paragraph(style="Certificate Text")
        self._run(p, "Submitted for the project viva-voce examination held on ______________________", size=14)
        for _ in range(3):
            d.add_paragraph()
        t2 = d.add_table(rows=1, cols=2)
        for c_i, txt in enumerate(("INTERNAL EXAMINER", "EXTERNAL EXAMINER")):
            c = t2.rows[0].cells[c_i]
            c.width = Cm(7.75)
            cp = c.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            self._run(cp, txt, bold=True, size=12)

    def lists_pages(self):
        d = self.doc
        p = d.add_paragraph(style="Front Title")
        p.add_run("TABLE OF CONTENTS")
        p = d.add_paragraph()
        add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
        tiny_mark(p)
        for title, label in (("LIST OF FIGURES", "Figure"), ("LIST OF TABLES", "Table")):
            d.add_paragraph(title, style="Heading 1")
            p = d.add_paragraph()
            add_field(p, f'TOC \\h \\z \\c "{label}"', f"Right-click and choose Update Field to build the list of {label.lower()}s.")
            tiny_mark(p)

    def docs_title_page(self):
        d, m = self.doc, self.meta
        sec = d.sections[0]
        sec.different_first_page_header_footer = True
        page_number_format(sec, "lowerRoman", 1)
        self.footer_page_number(sec)
        for _ in range(5):
            d.add_paragraph()
        for text, size, bold in ((m["title"], 22, True), ("", 12, False), (m["subtitle"], 16, False), ("", 12, False),
                                 (m["authors"], 14, True), (m["university"], 13, False), (m["month_year"], 13, False)):
            p = d.add_paragraph(style="Title Page")
            self._run(p, text, bold=bold, size=size)
        p = d.add_paragraph(style="Front Title")
        p.add_run("CONTENTS")
        p = d.add_paragraph()
        add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")

    # ------------------------------------------------------------ build
    def set_properties(self):
        cp = self.doc.core_properties
        cp.title = self.meta["title"]
        cp.author = self.meta.get("authors") or ", ".join(s.split("(")[0].strip() for s in self.meta["students"].split(";"))
        cp.subject = self.meta.get("subject", "UG project report, Chandigarh University")
        cp.keywords = self.meta.get("keywords", "")
        cp.comments = "Built from Markdown by report/src/build_report.py"


def read_meta(path):
    meta = {}
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        m = re.match(r"^([a-z_]+):\s*(.*)$", line)
        if m:
            meta[m.group(1)] = m.group(2).strip()
    return meta


def build_report():
    md = REPORT / "md"
    meta = read_meta(md / "00_meta.md")
    b = Builder(meta, meta.get("short_title", meta["title"]))
    files = [md / "01_front_matter.md"] + sorted(md.glob("[1-5]0_*.md")) + [md / "90_references.md"] + sorted(md.glob("9[5-9]_*.md"))
    texts = [f.read_text(encoding="utf-8") for f in files]
    b.collect_refs(texts)
    b.load_citations((md / "90_references.md").read_text(encoding="utf-8"))
    b.title_page()
    front = texts[0].split("# ABSTRACT", 1)
    ack, rest = front[0].replace("\\newpage", ""), "# ABSTRACT" + front[1]
    b.render_markdown(ack)
    b.lists_pages()
    # abstract in its own style (TNR 14, double spacing, per the format)
    abstract, after = rest.split("\\newpage", 1)
    for chunk in abstract.split("\n\n"):
        c = chunk.strip()
        if not c:
            continue
        if c.startswith("# "):
            b.doc.add_paragraph(c[2:].strip().upper(), style="Heading 1")
        else:
            p = b.doc.add_paragraph(style="Abstract Text")
            b.add_inline(p, " ".join(x.strip() for x in c.split("\n")), size=14)
    b.render_markdown(after)
    for f, t in zip(files[1:], texts[1:]):
        b._in_references = f.name.startswith("90_")
        b.render_markdown(t)
    b._in_references = False
    b.set_properties()
    out = REPORT / "Project_Report.docx"
    b.doc.save(out)
    print(f"wrote {out}  ({len(b.refs)} numbered items)")


def build_docs():
    docs = ROOT / "docs"
    meta = read_meta(docs / "00_meta.md")
    b = Builder(meta, meta.get("short_title", meta["title"]), docs_mode=True)
    files = sorted(f for f in docs.glob("[0-9][0-9]_*.md") if not f.name.startswith("00_"))
    texts = [f.read_text(encoding="utf-8") for f in files]
    b.collect_refs(texts)
    b.docs_title_page()
    for t in texts:
        b.render_markdown(t)
    b.set_properties()
    out = docs / "Project_Documentation.docx"
    b.doc.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--docs", action="store_true", help="build docs/Project_Documentation.docx instead")
    ap.add_argument("--rotate-wide", action="store_true",
                    help="turn very wide figures sideways on their portrait page (off by default)")
    a = ap.parse_args()
    Builder.rotate_default = a.rotate_wide
    sys.exit(build_docs() if a.docs else build_report())
